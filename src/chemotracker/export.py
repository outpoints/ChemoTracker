"""Care-team exports: a styled Excel workbook and flat CSV files.

The database is the source of truth; these files are generated on demand
(and alongside every backup). Planned doses that were never logged appear as
"Not logged" so gaps are visible to the team.
"""

from __future__ import annotations

import csv
import io
import sqlite3
from dataclasses import dataclass
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from . import medinfo, schedule, store
from . import symptoms as sx
from .db import get_settings

INK = "0F3B40"
TEAL = "11616A"
MUTED = "5B6B6D"
FILL_HEAD = PatternFill("solid", fgColor=TEAL)
FILL_BAND = PatternFill("solid", fgColor="EEF3F2")
FILL_TAKEN = PatternFill("solid", fgColor="DDEFE2")
FILL_SKIP = PatternFill("solid", fgColor="FBEBC8")
FILL_MISS = PatternFill("solid", fgColor="F7DCD6")
FILL_SEV = {1: PatternFill("solid", fgColor="FFF4C7"), 2: PatternFill("solid", fgColor="FBD9B3"),
            3: PatternFill("solid", fgColor="F4B8AE")}
THIN = Side(style="thin", color="C9D3D2")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

KIND_SHORT = {"scheduled": "Scheduled", "as_needed": "As needed", "clinic": "Clinic"}
STATUS_TEXT = {"taken": "Taken", "skipped": "Skipped", "not_logged": "Not logged"}


@dataclass
class DoseRow:
    when: datetime | None
    planned: datetime | None
    cycle: int | None
    day: int | None
    medicine: str
    strength: str
    dose: str
    kind: str
    status: str
    reason: str
    helped: str
    dose_feel: str
    side_effects: str
    notes: str


@dataclass
class SymptomRow:
    when: datetime
    cycle: int | None
    day: int | None
    symptom: str
    severity: str
    reading: str
    score: int
    notes: str


def _cycle_day(conn: sqlite3.Connection, d: date, cache: dict) -> tuple[int | None, int | None]:
    if d not in cache:
        c = schedule.cycle_for_date(conn, d)
        cache[d] = (c["number"], schedule.day_number(c, d)) if c else (None, None)
    return cache[d]


def dose_rows(conn: sqlite3.Connection, start: date, end: date, now: datetime) -> list[DoseRow]:
    cache: dict = {}
    rows: list[DoseRow] = []
    for r in store.doses_between(conn, start, end):
        when = schedule.parse_local(r["taken_at"])
        planned = schedule.parse_local(r["slot"]) if r["slot"] and len(r["slot"]) > 10 else None
        ref = (when or planned or datetime.fromisoformat(r["slot"])).date()
        cyc, day = _cycle_day(conn, ref, cache)
        rows.append(DoseRow(
            when=when or planned or datetime.combine(ref, datetime.min.time()),
            planned=planned, cycle=cyc, day=day, medicine=_med_name(r["med_name"], r["med_brand"]),
            strength=r["med_strength"] or "", dose=r["dose"] or r["med_dose"] or "", kind=KIND_SHORT.get(r["kind"], r["kind"]),
            status=STATUS_TEXT[r["status"]], reason=r["reason"] or "", helped=store.HELPED.get(r["helped"] or "", ""),
            dose_feel=store.DOSE_FEEL.get(r["dose_feel"] or "", ""), side_effects=r["side_effects"] or "",
            notes=r["notes"] or "",
        ))
    # Planned doses in the window that nobody logged.
    seen_cycles = {c["id"]: c for c in schedule.all_cycles(conn)}
    for c in seen_cycles.values():
        for s in schedule.slots(conn, c):
            if s.log is not None or not (start <= s.date <= end):
                continue
            if s.status(now) not in ("missed", "late", "not_logged"):
                continue
            rows.append(DoseRow(
                when=s.dt if s.time else datetime.combine(s.date, datetime.min.time()), planned=s.dt if s.time else None,
                cycle=c["number"], day=s.day, medicine=_med_name(s.plan["med_name"], s.plan["med_brand"]),
                strength=s.plan["med_strength"] or "", dose=s.dose, kind=KIND_SHORT[s.kind], status="Not logged",
                reason="", helped="", dose_feel="", side_effects="", notes="",
            ))
    rows.sort(key=lambda r: (r.when or datetime.min, r.medicine))
    return rows


def symptom_rows(conn: sqlite3.Connection, start: date, end: date, fever_f: float) -> list[SymptomRow]:
    cache: dict = {}
    out = []
    for r in schedule.symptoms_between(conn, start, end):
        sym = sx.lookup(r["symptom"])
        when = schedule.parse_local(r["logged_at"])
        assert when is not None
        cyc, day = _cycle_day(conn, when.date(), cache)
        out.append(SymptomRow(
            when=when, cycle=cyc, day=day, symptom=r["label"] if r["symptom"] == "other" else sym.name,
            severity=sx.SEVERITY_NAMES.get(r["severity"], "") if r["severity"] is not None else "",
            reading=sym.describe(None, r["value"]) if r["value"] is not None else "",
            score=schedule.symptom_score(r, fever_f), notes=r["notes"] or "",
        ))
    return out


def _med_name(name: str, brand: str | None) -> str:
    return f"{name} ({brand})" if brand else name


def _plan_text(p: sqlite3.Row) -> str:
    bits = [medinfo.KIND_NAMES.get(p["kind"], p["kind"]), schedule.days_text(p["start_day"], p["end_day"])]
    times = schedule.plan_times(p)
    if times:
        bits.append(", ".join(schedule.fmt_time(t) for t in times))
    if p["label"]:
        bits.append(p["label"])
    if p["from_cycle"] != 1 or p["to_cycle"] is not None:
        bits.append(f"cycles {p['from_cycle']}–{p['to_cycle'] or 'on'}")
    return " · ".join(bits)


# ---------------------------------------------------------------- CSV


def doses_csv(conn: sqlite3.Connection, start: date, end: date, now: datetime) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Date", "Time", "Planned time", "Cycle", "Day", "Medicine", "Strength", "Dose", "Type", "Status",
                "Taken for", "Helped", "Dose felt", "Side effects", "Notes"])
    for r in dose_rows(conn, start, end, now):
        w.writerow([
            r.when.date().isoformat() if r.when else "",
            r.when.strftime("%H:%M") if r.when and r.status == "Taken" and r.kind != "Clinic" else "",
            r.planned.strftime("%H:%M") if r.planned else "",
            r.cycle or "", "" if r.day is None else r.day, r.medicine, r.strength, r.dose, r.kind, r.status,
            r.reason, r.helped, r.dose_feel, r.side_effects, r.notes,
        ])
    return buf.getvalue().encode("utf-8-sig")


def symptoms_csv(conn: sqlite3.Connection, start: date, end: date, fever_f: float) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Date", "Time", "Cycle", "Day", "Symptom", "Severity", "Reading", "Notes"])
    for r in symptom_rows(conn, start, end, fever_f):
        w.writerow([r.when.date().isoformat(), r.when.strftime("%H:%M"), r.cycle or "",
                    "" if r.day is None else r.day, r.symptom, r.severity, r.reading, r.notes])
    return buf.getvalue().encode("utf-8-sig")


# ---------------------------------------------------------------- Excel


def _header(ws: Worksheet, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = FILL_HEAD
        c.alignment = CENTER
        c.border = BOX


def _widths(ws: Worksheet, widths: list[float]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _print_setup(ws: Worksheet, header_row: int) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.print_options.gridLines = False


def _title(ws: Worksheet, text: str, sub: str) -> None:
    ws["A1"] = text
    ws["A1"].font = Font(size=16, bold=True, color=INK)
    ws["A2"] = sub
    ws["A2"].font = Font(size=10, color=MUTED)


def _table(ws: Worksheet, name: str, header_row: int, last_row: int, last_col: int) -> None:
    if last_row <= header_row:
        return
    ref = f"A{header_row}:{get_column_letter(last_col)}{last_row}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight15", showRowStripes=True)
    ws.add_table(t)


def build_workbook(conn: sqlite3.Connection, start: date, end: date, now: datetime) -> bytes:
    cfg = get_settings(conn)
    fever = float(cfg.get("fever_f") or 100.4)
    who = cfg.get("patient_name") or ""
    sub = f"{who + ' · ' if who else ''}{schedule.fmt_day(start)} to {schedule.fmt_day(end)} · generated {now:%Y-%m-%d %H:%M}"
    doses = dose_rows(conn, start, end, now)
    syms = symptom_rows(conn, start, end, fever)

    wb = Workbook()
    _summary(wb.active, conn, cfg, sub, doses, syms, fever)
    _daily(wb.create_sheet("Daily overview"), sub, start, end, doses, syms, fever, conn)
    _dose_sheet(wb.create_sheet("Dose log"), sub, doses)
    _symptom_sheet(wb.create_sheet("Symptoms"), sub, syms)
    for c in schedule.all_cycles(conn):
        c_start, c_end = schedule.date_of(c, 0), schedule.date_of(c, schedule.cycle_span(conn, c))
        if c_end >= start and c_start <= end:
            _grid_sheet(wb.create_sheet(f"Cycle {c['number']} grid"[:31]), conn, c, now)
    _medicine_sheet(wb.create_sheet("Medicines"), conn, sub)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _summary(ws: Worksheet, conn, cfg, sub, doses: list[DoseRow], syms: list[SymptomRow], fever: float) -> None:
    ws.title = "Summary"
    _title(ws, "Supportive medicines & symptoms", sub)
    row = 4
    if cfg.get("clinic_name") or cfg.get("clinic_phone"):
        ws.cell(row=row, column=1, value=f"Care team: {cfg.get('clinic_name') or ''} {cfg.get('clinic_phone') or ''}".strip())
        ws.cell(row=row, column=1).font = Font(color=MUTED)
        row += 2
    ws.cell(row=row, column=1, value="Medicines in this period").font = Font(bold=True, size=12, color=INK)
    row += 1
    head = ["Medicine", "Type", "Planned", "Taken", "Skipped", "Not logged", "Helped: yes / a little / no",
            "Dose felt: right / too strong / not enough"]
    _header(ws, row, head)
    by_med: dict[tuple[str, str], list[DoseRow]] = {}
    for d in doses:
        by_med.setdefault((d.medicine, d.kind), []).append(d)
    first = row + 1
    for (name, kind), items in sorted(by_med.items(), key=lambda kv: (kv[0][1] != "Clinic", kv[0][1], kv[0][0])):
        row += 1
        taken = [i for i in items if i.status == "Taken"]
        values = [
            name, kind,
            "" if kind == "As needed" else len(items),
            len(taken),
            sum(1 for i in items if i.status == "Skipped"),
            sum(1 for i in items if i.status == "Not logged"),
            f"{sum(i.helped == 'Yes' for i in taken)} / {sum(i.helped == 'A little' for i in taken)} / {sum(i.helped == 'No' for i in taken)}",
            f"{sum(i.dose_feel == 'About right' for i in taken)} / {sum(i.dose_feel == 'Too strong' for i in taken)} / "
            f"{sum(i.dose_feel == 'Not enough' for i in taken)}",
        ]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = BOX
            c.alignment = WRAP if col == 1 else CENTER
        if (row - first) % 2:
            for col in range(1, len(head) + 1):
                ws.cell(row=row, column=col).fill = FILL_BAND
    if not by_med:
        row += 1
        ws.cell(row=row, column=1, value="No doses in this period.").font = Font(italic=True, color=MUTED)
    row += 2
    ws.cell(row=row, column=1, value="Symptoms: worst entry in this period").font = Font(bold=True, size=12, color=INK)
    row += 1
    _header(ws, row, ["Symptom", "Entries", "Worst", "When", "Cycle · day"])
    worst: dict[str, SymptomRow] = {}
    count: dict[str, int] = {}
    for s in syms:
        count[s.symptom] = count.get(s.symptom, 0) + 1
        if s.symptom not in worst or s.score > worst[s.symptom].score:
            worst[s.symptom] = s
    for name, s in sorted(worst.items(), key=lambda kv: -kv[1].score):
        row += 1
        values = [name, count[name], s.reading or s.severity, f"{s.when:%a %b %d, %I:%M %p}",
                  f"C{s.cycle} · D{s.day}" if s.cycle else ""]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = BOX
            c.alignment = WRAP if col == 1 else CENTER
        if s.score:
            ws.cell(row=row, column=3).fill = FILL_SEV[s.score]
    if not worst:
        row += 1
        ws.cell(row=row, column=1, value="No symptoms logged in this period.").font = Font(italic=True, color=MUTED)
    fevers = [s for s in syms if s.symptom == "Temperature" and s.reading and float(s.reading.split()[0]) >= fever]
    if fevers:
        row += 2
        ws.cell(row=row, column=1, value=f"Temperatures at or above {fever:g} °F").font = Font(bold=True, color="9B2C1F")
        for s in fevers:
            row += 1
            ws.cell(row=row, column=1, value=f"{s.when:%a %b %d, %I:%M %p}: {s.reading}")
    _widths(ws, [34, 12, 10, 10, 10, 12, 24, 30])
    ws.freeze_panes = None
    _print_setup(ws, 1)


def _daily(ws: Worksheet, sub, start: date, end: date, doses, syms, fever, conn) -> None:
    _title(ws, "Day by day", sub)
    head = ["Date", "Cycle", "Day", "Scheduled taken", "As-needed doses", "Nausea (worst)", "Vomiting (times)",
            "Highest temp °F", "Worst pain /10", "Other symptoms", "Notes"]
    _header(ws, 4, head)
    row = 4
    d = start
    cache: dict = {}
    while d <= end:
        day_doses = [x for x in doses if x.when and x.when.date() == d]
        day_syms = [s for s in syms if s.when.date() == d]
        if day_doses or day_syms:
            row += 1
            cyc, dn = _cycle_day(conn, d, cache)
            planned = [x for x in day_doses if x.kind != "As needed"]
            prn = [x for x in day_doses if x.kind == "As needed" and x.status == "Taken"]
            prn_text = ", ".join(sorted({x.medicine.split(" (")[0] for x in prn}))
            nausea = [s for s in day_syms if s.symptom == "Nausea"]
            vomit = [s for s in day_syms if s.symptom == "Vomiting"]
            temps = [float(s.reading.split()[0]) for s in day_syms if s.symptom == "Temperature" and s.reading]
            pains = [float(s.reading.split()[0]) for s in day_syms if s.symptom == "Pain" and s.reading]
            others = sorted({f"{s.symptom} ({s.severity or s.reading})" for s in day_syms
                             if s.symptom not in ("Nausea", "Vomiting", "Temperature", "Pain") and s.score})
            notes = "; ".join(s.notes for s in day_syms if s.notes)
            worst_nausea = max(nausea, key=lambda s: s.score, default=None)
            values = [
                d, cyc or "", "" if dn is None else dn,
                f"{sum(x.status == 'Taken' for x in planned)} of {len(planned)}" if planned else "",
                f"{len(prn)}: {prn_text}" if prn else "",
                worst_nausea.severity if worst_nausea else "",
                sum(int(s.reading.split()[0]) for s in vomit if s.reading) if vomit else "",
                max(temps) if temps else "",
                max(pains) if pains else "",
                ", ".join(others), notes,
            ]
            for col, v in enumerate(values, start=1):
                c = ws.cell(row=row, column=col, value=v)
                c.border = BOX
                c.alignment = WRAP if col in (5, 10, 11) else CENTER
            ws.cell(row=row, column=1).number_format = "ddd mmm d"
            if worst_nausea and worst_nausea.score:
                ws.cell(row=row, column=6).fill = FILL_SEV[worst_nausea.score]
            if temps and max(temps) >= fever:
                ws.cell(row=row, column=8).fill = FILL_SEV[3]
        d += timedelta(days=1)
    if row == 4:
        ws.cell(row=5, column=1, value="Nothing logged in this period.").font = Font(italic=True, color=MUTED)
    _widths(ws, [13, 7, 6, 11, 28, 12, 10, 10, 9, 34, 34])
    ws.freeze_panes = "B5"
    _print_setup(ws, 4)


def _dose_sheet(ws: Worksheet, sub, doses: list[DoseRow]) -> None:
    _title(ws, "Dose log", sub)
    head = ["Date", "Time taken", "Planned time", "Cycle", "Day", "Medicine", "Strength", "Dose", "Type", "Status",
            "Taken for", "Helped", "Dose felt", "Side effects", "Notes"]
    _header(ws, 4, head)
    row = 4
    for r in doses:
        row += 1
        values = [
            r.when.date() if r.when else None,
            r.when.strftime("%I:%M %p").lstrip("0") if r.when and r.status == "Taken" and r.kind != "Clinic" else "",
            r.planned.strftime("%I:%M %p").lstrip("0") if r.planned else "",
            r.cycle or "", "" if r.day is None else r.day, r.medicine, r.strength, r.dose, r.kind, r.status,
            r.reason, r.helped, r.dose_feel, r.side_effects, r.notes,
        ]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.alignment = WRAP if col in (6, 11, 14, 15) else CENTER
        ws.cell(row=row, column=1).number_format = "ddd mmm d"
        ws.cell(row=row, column=10).fill = {"Taken": FILL_TAKEN, "Skipped": FILL_SKIP}.get(r.status, FILL_MISS)
        if r.dose_feel in ("Too strong", "Not enough"):
            ws.cell(row=row, column=13).fill = FILL_SKIP
        if r.helped == "No":
            ws.cell(row=row, column=12).fill = FILL_SKIP
    _table(ws, "DoseLog", 4, row, len(head))
    _widths(ws, [13, 11, 11, 7, 6, 28, 13, 16, 11, 11, 16, 10, 12, 24, 30])
    ws.freeze_panes = "A5"
    _print_setup(ws, 4)


def _symptom_sheet(ws: Worksheet, sub, syms: list[SymptomRow]) -> None:
    _title(ws, "Symptoms", sub)
    head = ["Date", "Time", "Cycle", "Day", "Symptom", "Severity", "Reading", "Notes"]
    _header(ws, 4, head)
    row = 4
    for s in syms:
        row += 1
        values = [s.when.date(), s.when.strftime("%I:%M %p").lstrip("0"), s.cycle or "",
                  "" if s.day is None else s.day, s.symptom, s.severity, s.reading, s.notes]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.alignment = WRAP if col in (5, 8) else CENTER
        ws.cell(row=row, column=1).number_format = "ddd mmm d"
        if s.score:
            ws.cell(row=row, column=6 if s.severity else 7).fill = FILL_SEV[s.score]
    _table(ws, "Symptoms", 4, row, len(head))
    _widths(ws, [13, 10, 7, 6, 24, 11, 14, 44])
    ws.freeze_panes = "A5"
    _print_setup(ws, 4)


def _grid_sheet(ws: Worksheet, conn, cycle: sqlite3.Row, now: datetime) -> None:
    days, med_rows, sym_rows = schedule.grid(conn, cycle, now)
    regimen = f" · {cycle['regimen']}" if cycle["regimen"] else ""
    _title(ws, f"Cycle {cycle['number']}{regimen}", f"Day 1 (chemo) {schedule.fmt_day(schedule.day1_of(cycle))}. "
           "✓ taken (time), – skipped, ○ planned, ! not logged, ok = as-needed allowed. Symptoms show the worst entry each day.")
    head_row = 4
    ws.cell(row=head_row, column=1, value="Medicine")
    ws.cell(row=head_row, column=2, value="Type")
    for i, dn in enumerate(days):
        d = schedule.date_of(cycle, dn)
        ws.cell(row=head_row, column=3 + i, value=f"Day {dn}\n{d:%a %b} {d.day}")
    _header(ws, head_row, [str(ws.cell(row=head_row, column=c).value) for c in range(1, 3 + len(days))])
    ws.row_dimensions[head_row].height = 30
    row = head_row
    for r in med_rows:
        row += 1
        ws.cell(row=row, column=1, value=f"{r.name}\n{r.detail}" if r.detail else r.name).alignment = WRAP
        ws.cell(row=row, column=2, value=KIND_SHORT[r.kind]).alignment = CENTER
        for i, dn in enumerate(days):
            cell = r.cells[dn]
            parts = [f"✓ {schedule.fmt_time(t['taken_at'])}" if t["taken_at"] and r.kind != "clinic" else "✓"
                     for t in cell.taken]
            parts += ["–" for _ in cell.skipped]
            parts += ["!" for _ in range(cell.missing)]
            parts += ["○" for _ in range(cell.upcoming)]
            c = ws.cell(row=row, column=3 + i, value="\n".join(parts) or ("ok" if cell.allowed else None))
            c.alignment = CENTER
            c.border = BOX
            if cell.allowed and not parts:
                c.font = Font(color="9AA5A6", size=9)
                c.fill = FILL_BAND
            if cell.missing:
                c.fill = FILL_MISS
            elif cell.taken:
                c.fill = FILL_TAKEN
            elif cell.skipped:
                c.fill = FILL_SKIP
        ws.row_dimensions[row].height = 32
    if sym_rows:
        row += 1
        ws.cell(row=row, column=1, value="Symptoms").font = Font(bold=True, color=INK)
        fever = float(get_settings(conn).get("fever_f") or 100.4)
        for r in sym_rows:
            row += 1
            ws.cell(row=row, column=1, value=r.name).alignment = WRAP
            for i, dn in enumerate(days):
                entries = r.cells[dn]
                c = ws.cell(row=row, column=3 + i)
                c.border = BOX
                c.alignment = CENTER
                if not entries:
                    continue
                top = max(entries, key=lambda e: schedule.symptom_score(e, fever))
                sym = sx.lookup(top["symptom"])
                c.value = sym.describe(top["severity"], top["value"])
                score = schedule.symptom_score(top, fever)
                if score:
                    c.fill = FILL_SEV[score]
    _widths(ws, [30, 11] + [11] * len(days))
    ws.freeze_panes = ws.cell(row=head_row + 1, column=3)
    _print_setup(ws, head_row)


def _medicine_sheet(ws: Worksheet, conn, sub) -> None:
    _title(ws, "Medicines (from pharmacy labels)", sub)
    head = ["Medicine", "Strength", "Dose", "Label directions", "How it's used", "Min hours between", "Most per day",
            "Active", "Notes"]
    _header(ws, 4, head)
    row = 4
    for m in store.medications(conn, active=None):
        row += 1
        plan_text = "\n".join(_plan_text(p) for p in store.plans(conn, m["id"]))
        values = [_med_name(m["name"], m["brand"]), m["strength"] or "", m["dose"] or "", m["directions"] or "",
                  plan_text, m["min_hours_between"] or "", m["max_per_day"] or "", "Yes" if m["active"] else "No",
                  m["notes"] or ""]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.alignment = WRAP
            c.border = BOX
    _widths(ws, [28, 14, 16, 36, 40, 10, 9, 8, 30])
    ws.freeze_panes = "A5"
    _print_setup(ws, 4)
    ws.conditional_formatting.add(f"H5:H{max(row, 5)}", FormulaRule(formula=['$H5="No"'], font=Font(color="9AA5A6")))
